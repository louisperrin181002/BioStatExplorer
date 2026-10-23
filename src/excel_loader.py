# src/excel_loader.py
"""
Reads data tables from Excel (.xlsx/.xlsm) and CSV/TSV/TXT files.

* Merged cells (e.g. 'Experiment' or 'Strain' written once across a block
  of replicate rows) are filled into every row they cover.
* The header row is found automatically, so instrument preambles above the
  table (User:, Path:, Date: ...) are skipped.
* Empty rows/columns are dropped, duplicate column names made unique, and
  numbers stored as text converted.
Column names are kept exactly as written in the file.
"""
import csv
import io
import os
import re

import pandas as pd

TEXT_EXT = (".csv", ".tsv", ".txt")
EXCEL_EXT = (".xlsx", ".xlsm")
OVERFLOW = re.compile(r"^\s*(ovrflw|overflow|over|sat(urated)?|>.*|\*+)\s*$",
                      re.I)


class LoadReport:
    """What the loader did to a file, for the status bar / a message."""

    def __init__(self, name):
        self.name = name
        self.notes = []

    def __str__(self):
        return f"{self.name}: " + "; ".join(self.notes) if self.notes else ""


def sheet_names(path):
    """Sheet names of an Excel file ([] for text files)."""
    if not path.lower().endswith(EXCEL_EXT):
        return []
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def read_grid(path, sheet=None):
    """The raw cells of a sheet (or text file) as a DataFrame with no
    header, merged cells filled in. Used by the table and plate readers."""
    low = path.lower()
    if low.endswith(EXCEL_EXT):
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True)
        try:
            ws = wb[sheet] if sheet else wb.worksheets[0]
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            width = max((len(r) for r in rows), default=0)
            rows = [r + [None] * (width - len(r)) for r in rows]
            for rng in ws.merged_cells.ranges:
                r0, c0 = rng.min_row - 1, rng.min_col - 1
                if r0 >= len(rows) or c0 >= width:
                    continue
                v = rows[r0][c0]
                for r in range(r0, min(rng.max_row, len(rows))):
                    for c in range(c0, min(rng.max_col, width)):
                        rows[r][c] = v
            merged = len(ws.merged_cells.ranges)
        finally:
            wb.close()
        grid = pd.DataFrame(rows)
        grid.attrs["merged"] = merged
        return grid
    if low.endswith(".xls"):
        try:
            return pd.read_excel(path, sheet_name=sheet or 0, header=None)
        except ImportError as exc:
            raise ValueError("Old .xls files need the 'xlrd' package — or "
                             "re-save the file as .xlsx in Excel.") from exc
    raw = None
    for enc in ("utf-8-sig", "latin-1"):
        try:
            with open(path, encoding=enc, newline="") as f:
                raw = f.read()
            break
        except UnicodeDecodeError:
            continue
    try:
        dialect = csv.Sniffer().sniff(raw[:8192], delimiters=",;\t")
        sep = dialect.delimiter
    except csv.Error:
        sep = "\t" if low.endswith(".tsv") else ","
    return pd.read_csv(io.StringIO(raw), sep=sep, header=None, dtype=object,
                       skip_blank_lines=False, engine="python")


def _blank(v):
    return v is None or (isinstance(v, float) and pd.isna(v)) or \
        (isinstance(v, str) and not v.strip())


def _is_number(v):
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return not pd.isna(v)
    if isinstance(v, str):
        try:
            float(v.replace(",", "."))
            return True
        except ValueError:
            return False
    return False


def find_header_row(grid):
    """Index of the table's header row: the first row with at least two
    filled cells, mostly text, followed by a row of data."""
    n = len(grid)
    for r in range(min(n, 60)):
        row = grid.iloc[r].tolist()
        filled = [v for v in row if not _blank(v)]
        if len(filled) < 2:
            continue
        texts = [v for v in filled if not _is_number(v)]
        if len(texts) < max(2, 0.6 * len(filled)):
            continue
        if r + 1 >= n:
            return None
        nxt = [v for v in grid.iloc[r + 1].tolist() if not _blank(v)]
        if len(nxt) >= max(1, len(filled) // 2):
            return r
    return 0 if n else None


def table_from_grid(grid, report=None):
    """Turns a raw grid into a clean table (header detected, empty
    rows/columns dropped, numbers converted)."""
    report = report or LoadReport("")
    h = find_header_row(grid)
    if h is None:
        return pd.DataFrame()
    if h > 0:
        report.notes.append(f"table starts at row {h + 1}")
    header = ["" if _blank(v) else str(v).strip() for v in grid.iloc[h]]
    body = grid.iloc[h + 1:].reset_index(drop=True)
    keep = [i for i, name in enumerate(header)
            if name or body.iloc[:, i].map(lambda v: not _blank(v)).any()]
    body = body.iloc[:, keep]
    names, seen = [], {}
    for i in keep:
        name = header[i] or f"Column {i + 1}"
        if name in seen:
            seen[name] += 1
            name = f"{name} ({seen[name]})"
        else:
            seen[name] = 1
        names.append(name)
    body.columns = names
    body = body[~body.apply(lambda r: all(_blank(v) for v in r), axis=1)]
    body = body.reset_index(drop=True)
    for c in body.columns:
        col = body[c]
        if col.dtype == object:
            col = col.map(lambda v: v.strip() if isinstance(v, str) else v)
            col = col.map(lambda v: None if _blank(v) else v)
            body[c] = col
        vals = [v for v in body[c] if not _blank(v)]
        if not vals:
            continue
        nums = [v for v in vals if _is_number(v)]
        if len(nums) == len(vals) or (len(nums) >= 0.8 * len(vals)
                                      and len(vals) >= 5):
            bad = [v for v in vals if not _is_number(v)]
            conv = body[c].map(lambda v: float(v.replace(",", "."))
                               if isinstance(v, str) and _is_number(v) else v)
            body[c] = pd.to_numeric(conv, errors="coerce")
            if bad:
                over = sum(1 for v in bad if OVERFLOW.match(str(v)))
                report.notes.append(
                    f"'{c}': {len(bad)} non-numeric cell(s) left empty"
                    + (f" ({over} look like detector overflow)" if over else
                       f" (e.g. '{bad[0]}')"))
    return body


class DataLoader:
    def load_table(self, path, sheet=None):
        """(DataFrame, LoadReport) for one file/sheet."""
        report = LoadReport(os.path.basename(path) + (f" [{sheet}]"
                                                      if sheet else ""))
        grid = read_grid(path, sheet)
        if grid.attrs.get("merged"):
            report.notes.append(f"{grid.attrs['merged']} merged cell "
                                "range(s) filled into every row")
        return table_from_grid(grid, report), report

    def load_and_merge(self, file_paths, sheets=None):
        """
        Loads each file (sheets: {path: sheet name}), tags rows with their
        'Source File', and stacks them. Returns (DataFrame, numeric columns,
        other columns, [LoadReport]).
        """
        sheets = sheets or {}
        frames, reports = [], []
        for path in file_paths:
            if not os.path.exists(path):
                continue
            df, rep = self.load_table(path, sheets.get(path))
            if df.empty:
                rep.notes.append("no table found")
                reports.append(rep)
                continue
            df["Source File"] = os.path.basename(path)
            frames.append(df)
            reports.append(rep)
        if not frames:
            return pd.DataFrame(), [], [], reports
        merged = pd.concat(frames, ignore_index=True, sort=False)
        numeric = [c for c in merged.columns
                   if pd.api.types.is_numeric_dtype(merged[c])]
        other = [c for c in merged.columns if c not in numeric]
        return merged, numeric, other, reports
